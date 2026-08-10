#!/usr/bin/env python3
"""Write the ready-to-use text for one page: three .docx to read and approve,
three .html to paste into the page editor.

The URLs stay as they are — nothing is created or redirected — so each file is
simply the final content of a page that already exists, in one language, with
cross-references pointing at the live pages of that same language.

The HTML is deliberately plain: headings, paragraphs, lists, links. No classes,
no inline styles, nothing that fights the site's own stylesheet or breaks when
pasted into an editor.

Usage:  python3 build_pages.py <translations.json> <out_dir> [doc-key ...]
"""
import html
import json
import os
import re
import sys

LANGS = (("pl", "Polski"), ("ua", "Українська"), ("en", "English"))
LINK = re.compile(r"\[\[([^|]+)\|([^\]]+)\]\]")

# Numbered clause, bullet, or a heading-looking line.
NUMBERED = re.compile(r"^\s*(\d+)[.)]\s+(.*)")
BULLET = re.compile(r"^\s*[-–—•]\s+(.*)")
SECTION = re.compile(r"^\s*(§\s*\d+\.?|\d+\.)\s*[A-ZА-ЯЄІЇҐŁŚŻŹĆŃÓĄĘ]")


def strip_links(s):
    return LINK.sub(r"\1", s)


def esc_with_links(s):
    """Escape the text, then turn [[name|url]] markers into real anchors."""
    out, last = [], 0
    for m in LINK.finditer(s):
        out.append(html.escape(s[last:m.start()]))
        out.append(f'<a href="{html.escape(m.group(2))}">'
                   f'{html.escape(m.group(1))}</a>')
        last = m.end()
    out.append(html.escape(s[last:]))
    return "".join(out)


def classify(line):
    """What this paragraph is, so the markup carries its structure.

    The documents use "1." for two different things: a section title
    ("1. Prawo do odstąpienia od umowy") and a clause inside it ("1. Mają
    Państwo prawo odstąpić od umowy w terminie 14 dni..."). What separates them
    is shape, not the number — a title is short and does not end a sentence,
    a clause is a sentence and closes with a full stop.
    """
    if re.match(r"^\s*§\s*\d+", line):
        return "h2"

    m = NUMBERED.match(line)
    if m:
        return "h2" if _looks_like_title(line) else "li-num"
    if BULLET.match(line):
        return "li-bul"
    if _looks_like_title(line) and "·" not in line:
        return "h2"
    return "p"


def _looks_like_title(line):
    return (len(line) < 90
            and not line.endswith((".", ",", ";", ":", "!", "?")))


def build_html(doc, lang):
    title = doc.get(f"title_{lang}") or doc.get("title_pl") or doc["title"]
    parts = [f"<h1>{html.escape(title)}</h1>"]
    open_list = None

    def close():
        nonlocal open_list
        if open_list:
            parts.append(f"</{open_list}>")
            open_list = None

    for b in doc["blocks"]:
        line = (b.get(lang) or "").strip()
        if not line:
            continue
        # Role comes from the Polish source so all three languages keep the
        # same structure — a heading in one is a heading in all.
        kind = classify((b.get("pl") or "").strip() or line)
        if kind == "li-num":
            if open_list != "ol":
                close()
                parts.append("<ol>")
                open_list = "ol"
            m = NUMBERED.match(line)
            parts.append(f"  <li>{esc_with_links(m.group(2) if m else line)}</li>")
        elif kind == "li-bul":
            if open_list != "ul":
                close()
                parts.append("<ul>")
                open_list = "ul"
            m = BULLET.match(line)
            parts.append(f"  <li>{esc_with_links(m.group(1) if m else line)}</li>")
        elif kind == "h2":
            close()
            parts.append(f"<h2>{esc_with_links(line)}</h2>")
        else:
            close()
            parts.append(f"<p>{esc_with_links(line)}</p>")
    close()
    return "\n".join(parts) + "\n"


def build_docx(doc, lang, path):
    import docx
    from docx.shared import Pt

    d = docx.Document()
    style = d.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    title = doc.get(f"title_{lang}") or doc.get("title_pl") or doc["title"]
    d.add_heading(title, level=0)

    page = (doc.get("page") or {}).get(lang, "")
    if page:
        note = d.add_paragraph()
        run = note.add_run(f"Страница: {page}")
        run.italic = True
        run.font.size = Pt(9)

    for b in doc["blocks"]:
        line = (b.get(lang) or "").strip()
        if not line:
            continue
        kind = classify((b.get("pl") or "").strip() or line)
        text = strip_links(line)
        if kind == "h2":
            d.add_heading(text, level=1)
        elif kind == "li-num":
            m = NUMBERED.match(text)
            d.add_paragraph(m.group(2) if m else text, style="List Number")
        elif kind == "li-bul":
            m = BULLET.match(text)
            d.add_paragraph(m.group(1) if m else text, style="List Bullet")
        else:
            d.add_paragraph(text)

    # Links are lost in the plain paragraphs above, so list them once at the end
    # — the .docx is for reading and approval, the .html carries the live links.
    links = [l for l in (doc.get("links") or [])
             if l["url"].endswith(l["url"].split("/")[-1])]
    seen, rows = set(), []
    for b in doc["blocks"]:
        for name, u in LINK.findall(b.get(lang) or ""):
            if u in seen:
                continue
            seen.add(u)
            rows.append((name, u))
    if rows:
        d.add_page_break()
        d.add_heading("Ссылки в тексте", level=1)
        for name, u in rows:
            p = d.add_paragraph(style="List Bullet")
            p.add_run(f"{name} — ").bold = True
            p.add_run(u).font.size = Pt(9)

    d.save(path)


def main():
    tr_path, out_root = sys.argv[1], sys.argv[2]
    wanted = set(sys.argv[3:])
    data = json.load(open(tr_path, encoding="utf-8"))

    made = 0
    for doc in data["documents"]:
        if wanted and doc["doc"] not in wanted:
            continue
        slug = ((doc.get("page") or {}).get("pl", "")
                .rstrip("/").split("/")[-1]) or doc["doc"]
        folder = os.path.join(out_root, slug)
        os.makedirs(folder, exist_ok=True)

        for lang, label in LANGS:
            base = os.path.join(folder, f"{slug}.{lang}")
            open(base + ".html", "w", encoding="utf-8").write(build_html(doc, lang))
            build_docx(doc, lang, base + ".docx")
            made += 2

        print(f"  {slug:<42} {doc['title']}")
        for lang, _ in LANGS:
            page = (doc.get('page') or {}).get(lang, '')
            print(f"      {lang}  {slug}.{lang}.docx / .html   →  {page}")

    print(f"\nсоздано файлов: {made}")


if __name__ == "__main__":
    main()
